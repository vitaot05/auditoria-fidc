import os
import json
import holidays
import openpyxl
import pandas as pd

from datetime import datetime, timedelta

# ==========================================================
# CONFIG
# ==========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

ARQ_HISTORICO = os.path.join(BASE_DIR, "historico.json")
ARQ_DADOS = os.path.join(BASE_DIR, "dados.json")

USER = os.environ.get("USERPROFILE")

FUNDOS = [
    {
        "nome": "Ártico",
        "chave": "ARTICO",
        "pasta": f"{USER}\\OneDrive - BBS Capital\\BBS Capital Partners - Documentos\\01. BUs\\01. Struc Funds\\01. Credit\\01. Projetos\\Ártico FIDC\\Reporting\\2026\\Arquivos Base"
    },
    {
        "nome": "Alvorada",
        "chave": "ALVORADA",
        "pasta": f"{USER}\\OneDrive - BBS Capital\\BBS Capital Partners - Documentos\\01. BUs\\01. Struc Funds\\01. Credit\\01. Projetos\\Alvorada FIDC\\Reporting\\2026\\Arquivos Base"
    },
    {
        "nome": "Opera",
        "chave": "OPERA",
        "pasta": f"{USER}\\OneDrive - BBS Capital\\BBS Capital Partners - Documentos\\01. BUs\\01. Struc Funds\\01. Credit\\01. Projetos\\Opera FIDC\\Reporting\\2026\\Arquivos Base"
    }
]

# ==========================================================
# FERIADOS
# ==========================================================

ano_atual = datetime.now().year
FERIADOS_B3 = holidays.BR(years=ano_atual, subdiv='SP')

def calcula_pascoa(ano):
    a = ano % 19
    b = ano // 100
    c = ano % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return datetime(ano, mes, dia).date()

pascoa = calcula_pascoa(ano_atual)

FERIADOS_B3[pascoa - timedelta(days=48)] = "Carnaval"
FERIADOS_B3[pascoa - timedelta(days=47)] = "Carnaval"
FERIADOS_B3[pascoa + timedelta(days=60)] = "Corpus Christi"

def obter_dia_util_retroativo(n):
    d = datetime.now()
    encontrados = 0

    while encontrados < n:
        d -= timedelta(days=1)

        if d.weekday() < 5 and d.date() not in FERIADOS_B3:
            encontrados += 1

    return d

D1 = obter_dia_util_retroativo(1)
D2 = obter_dia_util_retroativo(2)

STR_NOME_D1 = D1.strftime("%d-%m")
STR_NOME_D2 = D2.strftime("%d-%m")

STR_PLANILHA_D1_BR = D1.strftime("%d/%m/%Y")
STR_PLANILHA_D1_ISO = D1.strftime("%Y-%m-%d")
STR_PLANILHA_D2_BR = D2.strftime("%d/%m/%Y")

# ==========================================================
# RESULTADOS
# ==========================================================

resultado_final = []

def registrar(fundo, arquivo, status, motivo=""):

    data_auditada = (
        STR_PLANILHA_D2_BR
        if arquivo == "Carteira"
        else STR_PLANILHA_D1_BR
    )

    resultado_final.append({
        "data_execucao": datetime.now().strftime("%Y-%m-%d"),
        "hora_execucao": datetime.now().strftime("%H:%M:%S"),
        "data_auditada": data_auditada,
        "fundo": fundo,
        "arquivo": arquivo,
        "status": status,
        "motivo": motivo
    })

# ==========================================================
# HELPERS
# ==========================================================

def normalizar(texto):

    t = str(texto).upper()

    substituicoes = {
        "Á": "A",
        "É": "E",
        "Í": "I",
        "Ó": "O",
        "Ú": "U"
    }

    for original, novo in substituicoes.items():
        t = t.replace(original, novo)

    return t

def localizar_arquivo_exigente(pasta, prefixo, data_esperada_str):

    if not os.path.exists(pasta):
        return None

    encontrados = []

    for nome in os.listdir(pasta):

        if not nome.lower().endswith(
            (".xlsx", ".xlsm", ".xls")
        ):
            continue

        if (
            normalizar(prefixo) in normalizar(nome)
            and
            data_esperada_str in nome
        ):
            encontrados.append(
                os.path.join(pasta, nome)
            )

    if not encontrados:
        return None

    return max(
        encontrados,
        key=os.path.getmtime
    )

# ==========================================================
# COLE SUAS 4 FUNÇÕES AQUI
# ==========================================================

def auditar_estoque(fundo):
    try:
        pasta = os.path.join(fundo["pasta"], "Estoques")
        arquivo = localizar_arquivo_exigente(pasta, "Estoque", STR_NOME_D1)

        if not arquivo:
            registrar(fundo["nome"], "Estoque", "ERRO", f"Arquivo com data {STR_NOME_D1} não encontrado")
            return

        print(f"    -> Lendo: {os.path.basename(arquivo)}")

        wb = openpyxl.load_workbook(arquivo, data_only=True)
        ws = wb.active

        if normalizar(str(ws["A1"].value)) != "NOME_FUNDO":
            raise Exception("A1 incorreto")

        if fundo["chave"] not in normalizar(str(ws["A2"].value)):
            raise Exception(f"A2 incorreto (esperado {fundo['chave']})")

        coluna_data = None
        for col in range(1, ws.max_column + 1):
            valor_cabecalho = str(ws.cell(row=1, column=col).value)
            if "DATA_REFERENCIA" in normalizar(valor_cabecalho).replace(" ", ""):
                coluna_data = col
                break

        if coluna_data is None:
            raise Exception("Cabeçalho 'DATA_REFERENCIA' não encontrado na Linha 1")

        data_planilha = str(ws.cell(row=2, column=coluna_data).value)

        if STR_PLANILHA_D1_BR not in data_planilha:
            raise Exception(f"Data incorreta ({data_planilha[:15]}). Esperado: {STR_PLANILHA_D1_BR}")

        registrar(fundo["nome"], "Estoque", "OK")
    except Exception as e:
        registrar(fundo["nome"], "Estoque", "ERRO", str(e))

def auditar_aquisicoes(fundo):
    try:
        pasta = os.path.join(fundo["pasta"], "Aquisições")
        arquivo = localizar_arquivo_exigente(pasta, "Aquisi", STR_NOME_D1)

        if not arquivo:
            registrar(fundo["nome"], "Aquisições", "ERRO", f"Arquivo com data {STR_NOME_D1} não encontrado")
            return

        print(f"    -> Lendo: {os.path.basename(arquivo)}")

        df = pd.read_excel(arquivo, nrows=2)

        if str(df.columns[0]).strip() != "FUNDO":
            raise Exception("Coluna FUNDO incorreta")

        nome_fundo_arq = normalizar(str(df.iloc[0, 0]))
        if fundo["chave"] not in nome_fundo_arq:
            raise Exception("Fundo incorreto")

        data_planilha = str(df.iloc[0, 2])
        if STR_PLANILHA_D1_ISO not in data_planilha:
            raise Exception(f"Data incorreta ({data_planilha[:15]}). Esperado: {STR_PLANILHA_D1_ISO}")

        registrar(fundo["nome"], "Aquisições", "OK")
    except Exception as e:
        registrar(fundo["nome"], "Aquisições", "ERRO", str(e))

def auditar_liquidados(fundo):
    try:
        pasta = os.path.join(fundo["pasta"], "Liquidados")
        arquivo = localizar_arquivo_exigente(pasta, "Liquidados", STR_NOME_D1)

        if not arquivo:
            registrar(fundo["nome"], "Liquidados", "ERRO", f"Arquivo com data {STR_NOME_D1} não encontrado")
            return

        print(f"    -> Lendo: {os.path.basename(arquivo)}")

        df = pd.read_excel(arquivo, nrows=2)

        if str(df.columns[0]).strip() != "FUNDO":
            raise Exception("FUNDO incorreto")

        nome_fundo_arq = normalizar(str(df.iloc[0, 0]))
        if fundo["chave"] not in nome_fundo_arq:
            raise Exception("Fundo incorreto")

        data_planilha = str(df.iloc[0, 1])
        if STR_PLANILHA_D1_BR not in data_planilha:
            raise Exception(f"Data incorreta ({data_planilha[:15]}). Esperado: {STR_PLANILHA_D1_BR}")

        registrar(fundo["nome"], "Liquidados", "OK")
    except Exception as e:
        registrar(fundo["nome"], "Liquidados", "ERRO", str(e))

def auditar_carteira(fundo):
    try:
        pasta = os.path.join(fundo["pasta"], "Carteiras")
        arquivo = localizar_arquivo_exigente(pasta, "Carteira", STR_NOME_D2)

        if not arquivo:
            registrar(fundo["nome"], "Carteira", "ERRO", f"Arquivo com data {STR_NOME_D2} não encontrado")
            return

        print(f"    -> Lendo: {os.path.basename(arquivo)}")

        data_esperada_planilha = f"{STR_PLANILHA_D2_BR} A {STR_PLANILHA_D2_BR}"

        wb = openpyxl.load_workbook(arquivo, data_only=True)
        ws = wb.active

        linha1 = normalizar(str(ws["A1"].value))
        linha2 = normalizar(str(ws["A2"].value))

        if "CLIENTE" not in linha1:
            raise Exception("A1 inválido")

        if fundo["chave"] not in linha1:
            raise Exception("Cliente inválido")

        if data_esperada_planilha not in linha2:
            raise Exception(f"Data incorreta ({linha2[:25]}). Esperado: {data_esperada_planilha}")

        registrar(fundo["nome"], "Carteira", "OK")
    except Exception as e:
        registrar(fundo["nome"], "Carteira", "ERRO", str(e))

# ==========================================================
# EXECUÇÃO
# ==========================================================

# ==========================================================
# NOVA LÓGICA DE LOOP (Substitua a sua seção de execução)
# ==========================================================

# 1. Função para encontrar o N-ésimo dia útil ANTES de uma data específica
def encontrar_dia_util_referencia(data_referencia, n):
    d = data_referencia
    encontrados = 0
    while encontrados < n:
        d -= timedelta(days=1)
        if d.weekday() < 5 and d.date() not in FERIADOS_B3:
            encontrados += 1
    return d

# 2. Definir o range: últimos 3 meses (~65 dias úteis)
dias_para_voltar = 65 
datas_analise = []

for i in range(1, dias_para_voltar + 1):
    # Calculamos o D1 e D2 para cada dia do passado
    ref_d1 = encontrar_dia_util_referencia(datetime.now(), i)
    ref_d2 = encontrar_dia_util_referencia(ref_d1, 1) # D2 é sempre um dia útil antes de D1
    
    datas_analise.append({
        "d1": ref_d1,
        "d2": ref_d2,
        "str_d1": ref_d1.strftime("%d-%m"),
        "str_d2": ref_d2.strftime("%d-%m"),
        "planilha_d1_br": ref_d1.strftime("%d/%m/%Y"),
        "planilha_d1_iso": ref_d1.strftime("%Y-%m-%d"),
        "planilha_d2_br": ref_d2.strftime("%d/%m/%Y")
    })

# 3. Loop principal
for data in datas_analise:
    # Atualizamos as variáveis globais que suas funções usam
    STR_NOME_D1 = data["str_d1"]
    STR_NOME_D2 = data["str_d2"]
    STR_PLANILHA_D1_BR = data["planilha_d1_br"]
    STR_PLANILHA_D1_ISO = data["planilha_d1_iso"]
    STR_PLANILHA_D2_BR = data["planilha_d2_br"]
    
    print(f"\n--- Auditando Data: {STR_PLANILHA_D1_BR} ---")
    
    for fundo in FUNDOS:
        print(f"  > Fundo: {fundo['nome']}")
        auditar_estoque(fundo)
        auditar_aquisicoes(fundo)
        auditar_liquidados(fundo)
        auditar_carteira(fundo)

# ==========================================================
# HISTÓRICO
# ==========================================================

historico = []

if os.path.exists(ARQ_HISTORICO):

    try:
        with open(
            ARQ_HISTORICO,
            "r",
            encoding="utf-8"
        ) as f:

            historico = json.load(f)

    except:
        historico = []

for novo in resultado_final:

    historico = [
        h for h in historico
        if not (
            h.get("data_auditada") == novo["data_auditada"]
            and h.get("fundo") == novo["fundo"]
            and h.get("arquivo") == novo["arquivo"]
        )
    ]

    historico.append(novo)

with open(
    ARQ_HISTORICO,
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        historico,
        f,
        ensure_ascii=False,
        indent=4
    )

# ==========================================================
# DASHBOARD
# ==========================================================

fundos_dashboard = {}

for item in historico:

    fundo = item["fundo"]

    if fundo not in fundos_dashboard:

        fundos_dashboard[fundo] = {
            "total": 0,
            "ok": 0,
            "erro": 0,
            "ultima_falha": None,
            "dias_auditados": set()
        }

    data_ref = item.get(
        "data_auditada",
        item.get("data_execucao", "")
    )

    fundos_dashboard[fundo]["dias_auditados"].add(
        data_ref
    )

    fundos_dashboard[fundo]["total"] += 1

    if item["status"] == "OK":

        fundos_dashboard[fundo]["ok"] += 1

    else:

        fundos_dashboard[fundo]["erro"] += 1

        fundos_dashboard[fundo]["ultima_falha"] = {
            "data": data_ref,
            "arquivo": item["arquivo"],
            "motivo": item.get("motivo", "")
        }

# ==========================================================
# KPIs GLOBAIS
# ==========================================================

total_fundos = len(fundos_dashboard)

total_registros = len(historico)

total_ok = sum(
    fundo["ok"]
    for fundo in fundos_dashboard.values()
)

total_erros = sum(
    fundo["erro"]
    for fundo in fundos_dashboard.values()
)

taxa_global = round(
    (total_ok / max(total_registros, 1)) * 100,
    2
)

# ==========================================================
# JSON FINAL
# ==========================================================

dados = {
    "ultima_atualizacao":
        datetime.now().strftime("%d/%m/%Y %H:%M:%S"),

    "kpis": {

        "fundos_monitorados":
            total_fundos,

        "taxa_global":
            taxa_global,

        "total_registros":
            total_registros,

        "total_falhas":
            total_erros

    },

    "fundos": []
}

for nome, info in fundos_dashboard.items():

    taxa = round(
        (
            info["ok"] /
            max(info["total"], 1)
        ) * 100,
        2
    )

    dados["fundos"].append({

        "nome":
            nome,

        "taxa_sucesso":
            taxa,

        "falhas":
            info["erro"],

        "total":
            info["total"],

        "dias_auditados":
            len(info["dias_auditados"]),

        "ultima_falha":
            info["ultima_falha"]

    })

# ==========================================================
# ORDENAÇÃO
# ==========================================================

dados["fundos"] = sorted(
    dados["fundos"],
    key=lambda x: x["nome"]
)

# ==========================================================
# SALVAR DADOS.JSON
# ==========================================================

with open(
    ARQ_DADOS,
    "w",
    encoding="utf-8"
) as f:

    json.dump(
        dados,
        f,
        ensure_ascii=False,
        indent=4
    )

print("\nDashboard atualizado.")

print(
    f"Fundos: {total_fundos} | "
    f"Registros: {total_registros} | "
    f"Falhas: {total_erros} | "
    f"Taxa Global: {taxa_global}%"
)