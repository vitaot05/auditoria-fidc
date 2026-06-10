import os
import json
import logging
import argparse
import subprocess
import unicodedata
import holidays
import openpyxl
import pandas as pd

from dataclasses import dataclass
from datetime import datetime, timedelta

# ==========================================================
# CONFIG
# ==========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

ARQ_HISTORICO = os.path.join(BASE_DIR, "historico.json")
ARQ_DADOS     = os.path.join(BASE_DIR, "dados.json")
ARQ_FUNDOS    = os.path.join(BASE_DIR, "fundos.json")

USER = os.environ.get("USERPROFILE", "")

# ==========================================================
# LOGGING
# ==========================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[logging.StreamHandler()]
)
log = logging.getLogger(__name__)

# ==========================================================
# FUNDOS (carregados de fundos.json)
# ==========================================================

try:
    with open(ARQ_FUNDOS, "r", encoding="utf-8") as f:
        FUNDOS = json.load(f)
    for fundo in FUNDOS:
        fundo["pasta"] = fundo["pasta"].replace("{USER}", USER)
    log.info(f"{len(FUNDOS)} fundo(s) carregado(s) de fundos.json")
except FileNotFoundError:
    log.error(f"fundos.json não encontrado em {BASE_DIR}")
    raise
except json.JSONDecodeError as e:
    log.error(f"fundos.json inválido: {e}")
    raise

# ==========================================================
# FERIADOS B3
# Usa apenas feriados nacionais (sem subdiv SP) para não incluir
# feriados municipais que a B3 não observa.
# Carnaval e Corpus Christi são calculados manualmente.
# ==========================================================

ano_atual = datetime.now().year
FERIADOS_B3 = holidays.BR(years=ano_atual)  # nacionais apenas

def calcula_pascoa(ano):
    a = ano % 19
    b = ano // 100
    c = ano % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    return datetime(ano, mes, dia).date()

pascoa = calcula_pascoa(ano_atual)
FERIADOS_B3[pascoa - timedelta(days=48)] = "Carnaval (segunda)"
FERIADOS_B3[pascoa - timedelta(days=47)] = "Carnaval (terça)"
FERIADOS_B3[pascoa + timedelta(days=60)] = "Corpus Christi"

def obter_dia_util_retroativo(n):
    d = datetime.now()
    encontrados = 0
    while encontrados < n:
        d -= timedelta(days=1)
        if d.weekday() < 5 and d.date() not in FERIADOS_B3:
            encontrados += 1
    return d


@dataclass
class Datas:
    d1: datetime
    d2: datetime
    nome_d1: str
    nome_d2: str
    planilha_d1_br: str
    planilha_d1_iso: str
    planilha_d2_br: str


def obter_datas(i: int) -> Datas:
    """Retorna o contexto de datas para o i-ésimo dia útil retroativo."""
    d1 = obter_dia_util_retroativo(i)
    d2 = obter_dia_util_retroativo(i + 1)
    return Datas(
        d1=d1,
        d2=d2,
        nome_d1=d1.strftime("%d-%m"),
        nome_d2=d2.strftime("%d-%m"),
        planilha_d1_br=d1.strftime("%d/%m/%Y"),
        planilha_d1_iso=d1.strftime("%Y-%m-%d"),
        planilha_d2_br=d2.strftime("%d/%m/%Y"),
    )

# ==========================================================
# RESULTADOS
# ==========================================================

resultado_final = []

def registrar(fundo, arquivo, status, datas: Datas, motivo=""):
    agora = datetime.now()
    data_auditada = datas.planilha_d2_br if arquivo == "Carteira" else datas.planilha_d1_br

    resultado_final.append({
        "data_execucao": agora.strftime("%Y-%m-%d"),
        "hora_execucao": agora.strftime("%H:%M:%S"),
        "data_auditada": data_auditada,
        "fundo":         fundo,
        "arquivo":       arquivo,
        "status":        status,
        "motivo":        motivo
    })

    if status == "OK":
        log.info(f"  [{fundo}] {arquivo} → OK")
    elif status == "AVISO":
        log.warning(f"  [{fundo}] {arquivo} → AVISO: {motivo}")
    else:
        log.warning(f"  [{fundo}] {arquivo} → ERRO: {motivo}")

# ==========================================================
# HELPERS
# ==========================================================

def normalizar(texto):
    t = unicodedata.normalize("NFKD", str(texto).upper())
    return "".join(c for c in t if not unicodedata.combining(c))

def eh_copia(nome):
    """Retorna True se o arquivo parece ser uma cópia (contém '- copia' no nome)."""
    return "COPIA" in normalizar(nome)

def localizar_arquivo_exigente(pasta, prefixo, data_esperada_str):
    if not os.path.exists(pasta):
        return None

    encontrados = []
    for nome in os.listdir(pasta):
        if not nome.lower().endswith((".xlsx", ".xlsm", ".xls")):
            continue
        if normalizar(prefixo) in normalizar(nome) and data_esperada_str in nome:
            encontrados.append(os.path.join(pasta, nome))

    if not encontrados:
        return None

    # Prefere arquivos sem "copia" no nome; desempata pelo mais recente
    encontrados.sort(key=lambda p: (eh_copia(os.path.basename(p)), -os.path.getmtime(p)))
    escolhido = encontrados[0]

    if len(encontrados) > 1:
        nomes = [os.path.basename(p) for p in encontrados]
        log.warning(
            f"Múltiplos arquivos encontrados para '{prefixo}' / '{data_esperada_str}' "
            f"em '{pasta}': {nomes}. Usando: {os.path.basename(escolhido)}"
        )

    return escolhido

def parse_data_flexivel(valor):
    """
    Converte valor para date. Aceita datetime/Timestamp ou strings nos formatos:
    DD/MM/YYYY, YYYY-MM-DD, DD-MM-YYYY, YYYY/MM/DD.
    Lança ValueError com mensagem detalhada se nenhum formato funcionar.
    """
    if hasattr(valor, "date"):
        return valor.date()

    texto = str(valor).strip()[:10]
    formatos = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"]

    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue

    raise ValueError(
        f"Data '{texto}' não reconhecida — "
        f"formatos tentados: {', '.join(formatos)}"
    )

# ==========================================================
# FUNÇÕES DE AUDITORIA
# ==========================================================

def auditar_estoque(fundo, datas: Datas):
    nome = fundo["nome"]
    try:
        pasta = os.path.join(fundo["pasta"], "Estoques")
        arquivo = localizar_arquivo_exigente(pasta, "Estoque", datas.nome_d1)

        if not arquivo:
            registrar(nome, "Estoque", "AVISO", datas,
                      f"Arquivo não encontrado para {datas.nome_d1} em '{pasta}'")
            return

        log.info(f"  [{nome}] Lendo Estoque: {os.path.basename(arquivo)}")
        wb = openpyxl.load_workbook(arquivo, data_only=True)
        ws = wb.active

        # A1 deve conter NOME_FUNDO
        a1 = normalizar(str(ws["A1"].value))
        if "NOME_FUNDO" not in a1:
            raise ValueError(
                f"A1 incorreto — esperado 'NOME_FUNDO', encontrado '{ws['A1'].value}'"
            )

        # A2 deve conter a chave do fundo
        a2 = normalizar(str(ws["A2"].value))
        if fundo["chave"] not in a2:
            raise ValueError(
                f"Fundo incorreto — esperado '{fundo['chave']}', "
                f"encontrado '{ws['A2'].value}'"
            )

        # Localizar coluna DATA_REFERENCIA
        coluna_data = None
        for col in range(1, ws.max_column + 1):
            cabecalho = normalizar(str(ws.cell(row=1, column=col).value)).replace(" ", "")
            if "DATA_REFERENCIA" in cabecalho:
                coluna_data = col
                break

        if coluna_data is None:
            raise ValueError("Coluna 'DATA_REFERENCIA' não encontrada na linha 1")

        val_data = ws.cell(row=2, column=coluna_data).value
        data_planilha = parse_data_flexivel(val_data)
        data_esperada = datas.d1.date()

        if data_planilha != data_esperada:
            raise ValueError(
                f"Data incorreta — esperado {datas.planilha_d1_br}, "
                f"encontrado {data_planilha.strftime('%d/%m/%Y')}"
            )

        registrar(nome, "Estoque", "OK", datas)

    except Exception as e:
        registrar(nome, "Estoque", "ERRO", datas, str(e))


def auditar_aquisicoes(fundo, datas: Datas):
    nome = fundo["nome"]
    try:
        pasta = os.path.join(fundo["pasta"], "Aquisições")
        arquivo = localizar_arquivo_exigente(pasta, "Aquisi", datas.nome_d1)

        if not arquivo:
            registrar(nome, "Aquisições", "AVISO", datas,
                      f"Arquivo não encontrado para {datas.nome_d1} em '{pasta}' "
                      f"(possível dia sem movimento)")
            return

        log.info(f"  [{nome}] Lendo Aquisições: {os.path.basename(arquivo)}")
        df = pd.read_excel(arquivo, nrows=2)

        # Primeira coluna deve ser FUNDO
        col0 = str(df.columns[0]).strip()
        if col0 != "FUNDO":
            raise ValueError(
                f"Primeira coluna incorreta — esperado 'FUNDO', encontrado '{col0}'"
            )

        # Verificar nome do fundo
        nome_fundo_arq = normalizar(str(df.iloc[0, 0]))
        if fundo["chave"] not in nome_fundo_arq:
            raise ValueError(
                f"Fundo incorreto — esperado '{fundo['chave']}', "
                f"encontrado '{df.iloc[0, 0]}'"
            )

        # Verificar data (coluna índice 2)
        val_data = df.iloc[0, 2]
        data_planilha = parse_data_flexivel(val_data)
        data_esperada = datas.d1.date()

        if data_planilha != data_esperada:
            raise ValueError(
                f"Data incorreta — esperado {datas.planilha_d1_br}, "
                f"encontrado {data_planilha.strftime('%d/%m/%Y')}"
            )

        registrar(nome, "Aquisições", "OK", datas)

    except Exception as e:
        registrar(nome, "Aquisições", "ERRO", datas, str(e))


def auditar_liquidados(fundo, datas: Datas):
    nome = fundo["nome"]
    try:
        pasta = os.path.join(fundo["pasta"], "Liquidados")
        arquivo = localizar_arquivo_exigente(pasta, "Liquidados", datas.nome_d1)

        if not arquivo:
            registrar(nome, "Liquidados", "AVISO", datas,
                      f"Arquivo não encontrado para {datas.nome_d1} em '{pasta}' "
                      f"(possível dia sem movimento)")
            return

        log.info(f"  [{nome}] Lendo Liquidados: {os.path.basename(arquivo)}")
        df = pd.read_excel(arquivo, nrows=2)

        # Primeira coluna deve ser FUNDO
        col0 = str(df.columns[0]).strip()
        if col0 != "FUNDO":
            raise ValueError(
                f"Primeira coluna incorreta — esperado 'FUNDO', encontrado '{col0}'"
            )

        # Verificar nome do fundo
        nome_fundo_arq = normalizar(str(df.iloc[0, 0]))
        if fundo["chave"] not in nome_fundo_arq:
            raise ValueError(
                f"Fundo incorreto — esperado '{fundo['chave']}', "
                f"encontrado '{df.iloc[0, 0]}'"
            )

        # Verificar data (coluna índice 1)
        val_data = df.iloc[0, 1]
        data_planilha = parse_data_flexivel(val_data)
        data_esperada = datas.d1.date()

        if data_planilha != data_esperada:
            raise ValueError(
                f"Data incorreta — esperado {datas.planilha_d1_br}, "
                f"encontrado {data_planilha.strftime('%d/%m/%Y')}"
            )

        registrar(nome, "Liquidados", "OK", datas)

    except Exception as e:
        registrar(nome, "Liquidados", "ERRO", datas, str(e))


def auditar_carteira(fundo, datas: Datas):
    nome = fundo["nome"]
    try:
        pasta = os.path.join(fundo["pasta"], "Carteiras")
        arquivo = localizar_arquivo_exigente(pasta, "Carteira", datas.nome_d2)

        if not arquivo:
            registrar(nome, "Carteira", "AVISO", datas,
                      f"Arquivo não encontrado para {datas.nome_d2} em '{pasta}'")
            return

        log.info(f"  [{nome}] Lendo Carteira: {os.path.basename(arquivo)}")
        wb = openpyxl.load_workbook(arquivo, data_only=True)
        ws = wb.active

        linha1 = normalizar(str(ws["A1"].value))
        linha2 = normalizar(str(ws["A2"].value))

        # A1 deve conter CLIENTE
        if "CLIENTE" not in linha1:
            raise ValueError(
                f"A1 incorreto — esperado conter 'CLIENTE', encontrado '{ws['A1'].value}'"
            )

        # A1 deve conter a chave do fundo
        if fundo["chave"] not in linha1:
            raise ValueError(
                f"Fundo incorreto — esperado '{fundo['chave']}' em A1, "
                f"encontrado '{ws['A1'].value}'"
            )

        # A2 deve conter o período no formato "DD/MM/YYYY A DD/MM/YYYY"
        data_esperada_planilha = normalizar(
            f"{datas.planilha_d2_br} A {datas.planilha_d2_br}"
        )
        if data_esperada_planilha not in linha2:
            raise ValueError(
                f"Data incorreta em A2 — esperado "
                f"'{datas.planilha_d2_br} A {datas.planilha_d2_br}', "
                f"encontrado '{ws['A2'].value}'"
            )

        registrar(nome, "Carteira", "OK", datas)

    except Exception as e:
        registrar(nome, "Carteira", "ERRO", datas, str(e))

# ==========================================================
# EXECUÇÃO PRINCIPAL
# ==========================================================

def main():
    parser = argparse.ArgumentParser(description="Auditoria de fundos FIDC")
    parser.add_argument(
        "--reprocessar", type=int, default=1, metavar="N",
        help="Número de dias úteis retroativos a auditar (padrão: 1)"
    )
    args = parser.parse_args()
    n = args.reprocessar

    log.info("=" * 60)
    if n == 1:
        log.info("Iniciando auditoria")
    else:
        log.info(f"Iniciando reprocessamento — {n} dias úteis retroativos")

    for i in range(1, n + 1):
        datas = obter_datas(i)
        if n > 1:
            log.info(f"\n── Dia {i}/{n}: D1 = {datas.planilha_d1_br}  |  D2 = {datas.planilha_d2_br}")
        else:
            log.info(f"D1 = {datas.planilha_d1_br}  |  D2 = {datas.planilha_d2_br}")

        for fundo in FUNDOS:
            log.info(f"\nAuditando {fundo['nome']}")
            auditar_estoque(fundo, datas)
            auditar_aquisicoes(fundo, datas)
            auditar_liquidados(fundo, datas)
            auditar_carteira(fundo, datas)

    # ==========================================================
    # HISTÓRICO
    # ==========================================================

    historico = []
    if os.path.exists(ARQ_HISTORICO):
        try:
            with open(ARQ_HISTORICO, "r", encoding="utf-8") as f:
                historico = json.load(f)
        except (json.JSONDecodeError, ValueError) as e:
            log.error(f"historico.json corrompido, iniciando do zero: {e}")
            historico = []

    for novo in resultado_final:
        historico = [
            h for h in historico
            if not (
                h.get("data_auditada") == novo["data_auditada"]
                and h.get("fundo")     == novo["fundo"]
                and h.get("arquivo")   == novo["arquivo"]
            )
        ]
        historico.append(novo)

    with open(ARQ_HISTORICO, "w", encoding="utf-8") as f:
        json.dump(historico, f, ensure_ascii=False, indent=4)

    # ==========================================================
    # DASHBOARD (dados.json)
    # ==========================================================

    fundos_dashboard = {}

    for item in historico:
        fundo = item["fundo"]
        if fundo not in fundos_dashboard:
            fundos_dashboard[fundo] = {
                "total": 0, "ok": 0, "erro": 0, "aviso": 0,
                "ultima_falha": None, "dias_auditados": set()
            }

        data_ref = item.get("data_auditada", item.get("data_execucao", ""))
        fundos_dashboard[fundo]["dias_auditados"].add(data_ref)
        fundos_dashboard[fundo]["total"] += 1

        status = item["status"]
        if status == "OK":
            fundos_dashboard[fundo]["ok"] += 1
        elif status == "AVISO":
            fundos_dashboard[fundo]["aviso"] += 1
        else:
            fundos_dashboard[fundo]["erro"] += 1
            fundos_dashboard[fundo]["ultima_falha"] = {
                "data":    data_ref,
                "arquivo": item["arquivo"],
                "motivo":  item.get("motivo", "")
            }

    total_fundos    = len(fundos_dashboard)
    total_registros = len(historico)
    total_ok        = sum(v["ok"]   for v in fundos_dashboard.values())
    total_erros     = sum(v["erro"] for v in fundos_dashboard.values())
    taxa_global     = round((total_ok / max(total_registros, 1)) * 100, 2)

    dados = {
        "ultima_atualizacao": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "kpis": {
            "fundos_monitorados": total_fundos,
            "taxa_global":        taxa_global,
            "total_registros":    total_registros,
            "total_falhas":       total_erros
        },
        "fundos": []
    }

    for nome, info in fundos_dashboard.items():
        taxa = round((info["ok"] / max(info["total"], 1)) * 100, 2)
        dados["fundos"].append({
            "nome":           nome,
            "taxa_sucesso":   taxa,
            "falhas":         info["erro"],
            "avisos":         info["aviso"],
            "total":          info["total"],
            "dias_auditados": len(info["dias_auditados"]),
            "ultima_falha":   info["ultima_falha"]
        })

    dados["fundos"] = sorted(dados["fundos"], key=lambda x: x["nome"])

    with open(ARQ_DADOS, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=4)

    log.info(
        f"\nDashboard atualizado — Fundos: {total_fundos} | "
        f"Registros: {total_registros} | Falhas: {total_erros} | "
        f"Taxa Global: {taxa_global}%"
    )

    # ==========================================================
    # GIT PUSH (apenas na execução diária normal)
    # ==========================================================

    if n > 1:
        log.warning(
            f"Modo reprocessamento ({n} dias): git push omitido. "
            "Faça o push manualmente quando necessário."
        )
    else:
        try:
            subprocess.run(["git", "pull", "--rebase"], cwd=BASE_DIR, check=True)
            subprocess.run(
                ["git", "add", "dados.json", "historico.json"],
                cwd=BASE_DIR, check=True
            )
            subprocess.run(
                ["git", "commit", "-m",
                 f"auditoria {datetime.now().strftime('%Y-%m-%d %H:%M')}"],
                cwd=BASE_DIR, check=True
            )
            subprocess.run(["git", "push"], cwd=BASE_DIR, check=True)
            log.info("Git push realizado com sucesso.")
        except subprocess.CalledProcessError as e:
            log.error(f"Git falhou: {e}")


if __name__ == "__main__":
    main()
